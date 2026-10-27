from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path


@dataclass(frozen=True)
class TextChunk:
    text: str
    page_number: int | None = None
    paragraph_number: int | None = None


def _clean(text: str) -> str:
    return " ".join(text.replace("\u00a0", " ").replace("\u2028", " ").split())


def read_docx(path: Path) -> list[TextChunk]:
    from docx import Document

    document = Document(path)
    chunks: list[TextChunk] = []
    index = 0

    def add(value: str) -> None:
        nonlocal index
        value = _clean(value)
        if value:
            index += 1
            chunks.append(TextChunk(value, paragraph_number=index))

    for section in document.sections:
        for paragraph in section.header.paragraphs:
            add(paragraph.text)
    for paragraph in document.paragraphs:
        add(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            add(" | ".join(_clean(cell.text) for cell in row.cells if _clean(cell.text)))
    for section in document.sections:
        for paragraph in section.footer.paragraphs:
            add(paragraph.text)

    if not chunks:
        raise ValueError("В DOCX не найден извлекаемый текст")
    return chunks


def read_pdf(path: Path) -> list[TextChunk]:
    from pypdf import PdfReader

    reader = PdfReader(path)
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:  # pragma: no cover - library-specific failure
            raise ValueError("PDF защищён паролем") from exc

    chunks: list[TextChunk] = []
    paragraph = 0
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        for line in text.splitlines():
            line = _clean(line)
            if line:
                paragraph += 1
                chunks.append(TextChunk(line, page_number, paragraph))
    if not chunks:
        raise ValueError("В PDF не найден текстовый слой")
    return chunks


def read_document(path: Path, file_format: str) -> list[TextChunk]:
    if file_format == "DOCX":
        return read_docx(path)
    if file_format == "PDF":
        return read_pdf(path)
    raise ValueError("Поддерживаются только PDF и DOCX")

