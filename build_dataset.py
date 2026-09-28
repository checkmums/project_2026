from __future__ import annotations

import json
import random
from datetime import date, timedelta
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parent
DATASET_DIR = ROOT / "contract_dataset_v1"
DOCUMENTS_DIR = DATASET_DIR / "documents"
GROUND_TRUTH_PATH = DATASET_DIR / "ground_truth.json"
FONT_REGULAR = "/System/Library/Fonts/Supplemental/Arial.ttf"
FONT_BOLD = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"


CUSTOMERS = [
    "ООО «Сибирские решения»",
    "АО «Вектор»",
    "ООО «Альфа-Трейд»",
    "ООО «Новые технологии»",
    "АО «Городские системы»",
    "ООО «Логистик Плюс»",
    "ООО «РегионСтрой»",
    "АО «Промресурс»",
    "ООО «Континент»",
    "ООО «Северный проект»",
]

CONTRACTORS = [
    "ООО «ТехноСервис»",
    "ИП Иванов Иван Иванович",
    "ООО «Контур»",
    "ООО «ПрофКонсалт»",
    "ИП Петрова Анна Сергеевна",
    "ООО «МастерГрупп»",
    "ООО «Деловые линии Сибири»",
    "ИП Соколов Максим Олегович",
    "ООО «ИнфоПроект»",
    "ООО «СтройКомплект»",
]

CONTRACT_TYPES = [
    ("Договор оказания услуг", "оказать консультационные и информационные услуги"),
    ("Договор поставки", "поставить оборудование согласно спецификации"),
    ("Договор подряда", "выполнить ремонтные и монтажные работы"),
    ("Договор аренды", "предоставить во временное пользование офисное помещение"),
    ("Договор технического обслуживания", "выполнить техническое обслуживание оборудования"),
]

AMOUNTS = [
    75_000,
    98_500,
    120_000,
    145_300,
    180_000,
    215_750,
    248_500,
    310_000,
    425_600,
    590_000,
    675_250,
    810_000,
    950_400,
    1_125_000,
    1_480_500,
    1_750_000,
    2_040_300,
    2_350_000,
    2_780_900,
    3_120_000,
]


def format_amount(amount: int, variant: int) -> str:
    spaced = f"{amount:,}".replace(",", " ")
    if variant == 0:
        return f"{spaced},00 руб."
    if variant == 1:
        return f"{spaced} рублей"
    if variant == 2:
        return f"{spaced}.00 RUB"
    return f"{spaced} (НДС не облагается) руб."


def format_date(value: date, variant: int) -> str:
    months = [
        "января", "февраля", "марта", "апреля", "мая", "июня",
        "июля", "августа", "сентября", "октября", "ноября", "декабря",
    ]
    if variant == 0:
        return value.strftime("%d.%m.%Y")
    if variant == 1:
        return f"«{value.day:02d}» {months[value.month - 1]} {value.year} г."
    if variant == 2:
        return value.strftime("%d/%m/%Y")
    return value.strftime("%d-%m-%Y")


def make_records() -> list[dict]:
    random.seed(20260922)
    start = date(2026, 1, 15)
    records = []
    for index in range(20):
        variant = index % 4
        extension = "docx" if index % 2 == 0 else "pdf"
        contract_type, obligation = CONTRACT_TYPES[index % len(CONTRACT_TYPES)]
        contract_date = start + timedelta(days=index * 11)
        contract_number = [
            f"{101 + index}/26",
            f"УС-{2026}-{index + 1:03d}",
            f"П-{index + 1:02d}.{contract_date.month:02d}",
            f"{500 + index}-А",
        ][variant]
        record = {
            "document_id": f"DOC-{index + 1:03d}",
            "filename": f"contract_{index + 1:02d}.{extension}",
            "file_format": extension.upper(),
            "contract_type": contract_type,
            "contract_number": contract_number,
            "contract_date": contract_date.isoformat(),
            "contract_date_in_text": format_date(contract_date, variant),
            "customer": CUSTOMERS[index % len(CUSTOMERS)],
            "contractor": CONTRACTORS[(index * 3 + 1) % len(CONTRACTORS)],
            "amount_rub": AMOUNTS[index],
            "amount_in_text": format_amount(AMOUNTS[index], variant),
            "template_variant": f"T{variant + 1}",
            "difficulty": ["easy", "medium", "medium", "hard"][variant],
            "is_synthetic": True,
            "obligation": obligation,
        }
        records.append(record)
    return records


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_text(cell, value: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(value)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(10)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_docx_footer(document: Document) -> None:
    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Учебный синтетический документ. Не имеет юридической силы.")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = None


def build_docx(record: dict, path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08

    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(record["contract_type"].upper())
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(15)

    variant = int(record["template_variant"][1:])
    if variant == 1:
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(f"№ {record['contract_number']} от {record['contract_date_in_text']}").bold = True
        details = [
            ("Заказчик", record["customer"]),
            ("Исполнитель", record["contractor"]),
            ("Сумма договора", record["amount_in_text"]),
        ]
        table = document.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for label, value in details:
            cells = table.add_row().cells
            set_cell_text(cells[0], label, True)
            set_cell_text(cells[1], value)
            shade_cell(cells[0], "E7EDF3")
    elif variant == 2:
        document.add_paragraph(
            f"г. Новосибирск                                      {record['contract_date_in_text']}"
        )
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.add_run(record["customer"]).bold = True
        paragraph.add_run(", именуемое в дальнейшем «Заказчик», и ")
        paragraph.add_run(record["contractor"]).bold = True
        paragraph.add_run(", именуемое в дальнейшем «Исполнитель», заключили настоящий договор ")
        paragraph.add_run(f"№ {record['contract_number']}").bold = True
        paragraph.add_run(" о нижеследующем.")
    elif variant == 3:
        table = document.add_table(rows=3, cols=2)
        table.style = "Table Grid"
        rows = [
            ("Номер документа", record["contract_number"]),
            ("Дата заключения", record["contract_date_in_text"]),
            ("Вид договора", record["contract_type"]),
        ]
        for cells, (label, value) in zip(table.rows, rows):
            set_cell_text(cells.cells[0], label, True)
            set_cell_text(cells.cells[1], value)
            shade_cell(cells.cells[0], "E7EDF3")
        document.add_paragraph(
            f"Сторона 1 (Заказчик): {record['customer']}. "
            f"Сторона 2 (Исполнитель): {record['contractor']}."
        )
    else:
        document.add_paragraph(f"г. Новосибирск, {record['contract_date_in_text']}")
        document.add_paragraph(f"Договор № {record['contract_number']}")
        document.add_paragraph(f"ИСПОЛНИТЕЛЬ: {record['contractor']}")
        document.add_paragraph(f"ЗАКАЗЧИК: {record['customer']}")

    document.add_heading("1 Предмет договора", level=1)
    document.add_paragraph(
        f"Исполнитель обязуется {record['obligation']}, а Заказчик обязуется принять результат и оплатить его на условиях настоящего договора."
    )
    document.add_heading("2 Стоимость и порядок расчетов", level=1)
    if variant == 2:
        document.add_paragraph(
            f"2.1 Общая стоимость работ по настоящему договору составляет {record['amount_in_text']}."
        )
    elif variant == 3:
        document.add_paragraph(
            f"Стоимость обязательств сторон согласована в размере {record['amount_in_text']}. Оплата производится после подписания акта."
        )
    else:
        document.add_paragraph(f"Сумма договора: {record['amount_in_text']}")
    document.add_heading("3 Срок действия", level=1)
    document.add_paragraph(
        "Договор вступает в силу с даты подписания и действует до полного исполнения сторонами обязательств."
    )
    document.add_heading("4 Подписи сторон", level=1)
    signatures = document.add_table(rows=2, cols=2)
    signatures.style = "Table Grid"
    set_cell_text(signatures.cell(0, 0), "Заказчик", True)
    set_cell_text(signatures.cell(0, 1), "Исполнитель", True)
    set_cell_text(signatures.cell(1, 0), record["customer"] + "\n____________ / подпись")
    set_cell_text(signatures.cell(1, 1), record["contractor"] + "\n____________ / подпись")
    shade_cell(signatures.cell(0, 0), "E7EDF3")
    shade_cell(signatures.cell(0, 1), "E7EDF3")
    add_docx_footer(document)
    document.save(path)


def pdf_styles():
    pdfmetrics.registerFont(TTFont("ArialDataset", FONT_REGULAR))
    pdfmetrics.registerFont(TTFont("ArialDatasetBold", FONT_BOLD))
    styles = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "DatasetTitle", parent=styles["Title"], fontName="ArialDatasetBold",
            fontSize=15, leading=18, alignment=TA_CENTER, spaceAfter=10,
        ),
        "body": ParagraphStyle(
            "DatasetBody", parent=styles["BodyText"], fontName="ArialDataset",
            fontSize=10.5, leading=14, alignment=TA_JUSTIFY, spaceAfter=7,
        ),
        "right": ParagraphStyle(
            "DatasetRight", parent=styles["BodyText"], fontName="ArialDataset",
            fontSize=10.5, leading=14, alignment=TA_RIGHT, spaceAfter=7,
        ),
        "heading": ParagraphStyle(
            "DatasetHeading", parent=styles["Heading2"], fontName="ArialDatasetBold",
            fontSize=11.5, leading=14, spaceBefore=7, spaceAfter=5,
        ),
        "small": ParagraphStyle(
            "DatasetSmall", parent=styles["BodyText"], fontName="ArialDataset",
            fontSize=8, leading=10, alignment=TA_CENTER, textColor=colors.HexColor("#666666"),
        ),
    }


def build_pdf(record: dict, path: Path, styles: dict) -> None:
    document = SimpleDocTemplate(
        str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
        topMargin=16 * mm, bottomMargin=15 * mm,
        title=record["contract_type"], author="Учебный проект",
    )
    story = [Paragraph(record["contract_type"].upper(), styles["title"])]
    variant = int(record["template_variant"][1:])
    if variant == 1:
        story.append(Paragraph(
            f"<b>№ {record['contract_number']} от {record['contract_date_in_text']}</b>", styles["title"]
        ))
        data = [
            ["Заказчик", record["customer"]],
            ["Исполнитель", record["contractor"]],
            ["Сумма договора", record["amount_in_text"]],
        ]
        table = Table(data, colWidths=[45 * mm, 115 * mm])
        table.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), "ArialDataset"),
            ("FONTNAME", (0, 0), (0, -1), "ArialDatasetBold"),
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E7EDF3")),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#B8C1CA")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 7),
            ("RIGHTPADDING", (0, 0), (-1, -1), 7),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        story.extend([table, Spacer(1, 7)])
    elif variant == 2:
        story.append(Paragraph(
            f"г. Новосибирск&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;{record['contract_date_in_text']}", styles["right"]
        ))
        story.append(Paragraph(
            f"<b>{record['customer']}</b>, именуемое в дальнейшем «Заказчик», и "
            f"<b>{record['contractor']}</b>, именуемое в дальнейшем «Исполнитель», "
            f"заключили настоящий договор <b>№ {record['contract_number']}</b> о нижеследующем.",
            styles["body"],
        ))
    elif variant == 3:
        data = [
            ["Номер документа", record["contract_number"]],
            ["Дата заключения", record["contract_date_in_text"]],
            ["Вид договора", record["contract_type"]],
        ]
        table = Table(data, colWidths=[48 * mm, 112 * mm])
        table.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), "ArialDataset"),
            ("FONTNAME", (0, 0), (0, -1), "ArialDatasetBold"),
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E7EDF3")),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#B8C1CA")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 7),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        story.extend([
            table,
            Spacer(1, 8),
            Paragraph(
                f"Сторона 1 (Заказчик): {record['customer']}. "
                f"Сторона 2 (Исполнитель): {record['contractor']}.", styles["body"]
            ),
        ])
    else:
        story.extend([
            Paragraph(f"г. Новосибирск, {record['contract_date_in_text']}", styles["body"]),
            Paragraph(f"Договор № {record['contract_number']}", styles["body"]),
            Paragraph(f"ИСПОЛНИТЕЛЬ: {record['contractor']}", styles["body"]),
            Paragraph(f"ЗАКАЗЧИК: {record['customer']}", styles["body"]),
        ])

    story.extend([
        Paragraph("1 Предмет договора", styles["heading"]),
        Paragraph(
            f"Исполнитель обязуется {record['obligation']}, а Заказчик обязуется принять результат и оплатить его на условиях настоящего договора.",
            styles["body"],
        ),
        Paragraph("2 Стоимость и порядок расчетов", styles["heading"]),
    ])
    if variant == 2:
        story.append(Paragraph(
            f"2.1 Общая стоимость работ по настоящему договору составляет {record['amount_in_text']}.", styles["body"]
        ))
    elif variant == 3:
        story.append(Paragraph(
            f"Стоимость обязательств сторон согласована в размере {record['amount_in_text']}. Оплата производится после подписания акта.",
            styles["body"],
        ))
    else:
        story.append(Paragraph(f"Сумма договора: {record['amount_in_text']}", styles["body"]))
    story.extend([
        Paragraph("3 Срок действия", styles["heading"]),
        Paragraph(
            "Договор вступает в силу с даты подписания и действует до полного исполнения сторонами обязательств.", styles["body"]
        ),
        Paragraph("4 Подписи сторон", styles["heading"]),
    ])
    signatures = Table(
        [
            ["Заказчик", "Исполнитель"],
            [record["customer"] + "\n\n____________ / подпись", record["contractor"] + "\n\n____________ / подпись"],
        ],
        colWidths=[80 * mm, 80 * mm],
    )
    signatures.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "ArialDataset"),
        ("FONTNAME", (0, 0), (-1, 0), "ArialDatasetBold"),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E7EDF3")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#B8C1CA")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 1), (-1, 1), 18),
    ]))
    story.extend([
        signatures,
        Spacer(1, 14),
        Paragraph("Учебный синтетический документ. Не имеет юридической силы.", styles["small"]),
    ])
    document.build(story)


def build_architecture_svg(path: Path) -> None:
    svg = """<svg xmlns="http://www.w3.org/2000/svg" width="1400" height="430" viewBox="0 0 1400 430">
<rect width="1400" height="430" fill="#ffffff"/>
<text x="60" y="62" font-family="Arial" font-size="30" font-weight="700" fill="#17212b">Архитектура прототипа обработки договоров</text>
<text x="60" y="96" font-family="Arial" font-size="17" fill="#59636e">Первая версия: локальный Python-конвейер без отдельного backend и сложного интерфейса</text>
<defs><marker id="arrow" markerWidth="10" markerHeight="10" refX="8" refY="3" orient="auto"><path d="M0,0 L0,6 L9,3 z" fill="#66788a"/></marker></defs>
<g font-family="Arial" text-anchor="middle">
  <g><rect x="55" y="155" width="190" height="105" rx="10" fill="#eef3f7" stroke="#66819a" stroke-width="2"/><text x="150" y="194" font-size="20" font-weight="700" fill="#203040">Документы</text><text x="150" y="224" font-size="16" fill="#465866">PDF и DOCX</text></g>
  <g><rect x="295" y="155" width="210" height="105" rx="10" fill="#eef3f7" stroke="#66819a" stroke-width="2"/><text x="400" y="190" font-size="20" font-weight="700" fill="#203040">Чтение текста</text><text x="400" y="220" font-size="15" fill="#465866">pypdf</text><text x="400" y="242" font-size="15" fill="#465866">python-docx</text></g>
  <g><rect x="555" y="155" width="230" height="105" rx="10" fill="#fff5dd" stroke="#b18a35" stroke-width="2"/><text x="670" y="190" font-size="20" font-weight="700" fill="#4a3a17">Извлечение полей</text><text x="670" y="220" font-size="15" fill="#665527">RegEx</text><text x="670" y="242" font-size="15" fill="#665527">Natasha позднее</text></g>
  <g><rect x="835" y="155" width="220" height="105" rx="10" fill="#eef7ef" stroke="#5e8b65" stroke-width="2"/><text x="945" y="190" font-size="20" font-weight="700" fill="#243e28">Проверка</text><text x="945" y="220" font-size="15" fill="#49644d">формат даты</text><text x="945" y="242" font-size="15" fill="#49644d">сумма и стороны</text></g>
  <g><rect x="1105" y="155" width="230" height="105" rx="10" fill="#eef3f7" stroke="#66819a" stroke-width="2"/><text x="1220" y="190" font-size="20" font-weight="700" fill="#203040">Результат</text><text x="1220" y="220" font-size="15" fill="#465866">CSV на первом этапе</text><text x="1220" y="242" font-size="15" fill="#465866">SQLite позднее</text></g>
</g>
<g stroke="#66788a" stroke-width="3" fill="none" marker-end="url(#arrow)"><path d="M245 207 H285"/><path d="M505 207 H545"/><path d="M785 207 H825"/><path d="M1055 207 H1095"/></g>
<rect x="380" y="326" width="640" height="62" rx="8" fill="#f7f7f7" stroke="#b5bcc3"/>
<text x="700" y="352" text-anchor="middle" font-family="Arial" font-size="16" font-weight="700" fill="#39434c">Эталонная разметка датасета</text>
<text x="700" y="376" text-anchor="middle" font-family="Arial" font-size="15" fill="#5b6570">нужна для сравнения ожидаемых и извлеченных значений</text>
<path d="M700 326 V273" stroke="#87939e" stroke-width="2" stroke-dasharray="7 6" marker-end="url(#arrow)"/>
</svg>"""
    path.write_text(svg, encoding="utf-8")


def main() -> None:
    DOCUMENTS_DIR.mkdir(parents=True, exist_ok=True)
    records = make_records()
    styles = pdf_styles()
    for record in records:
        output = DOCUMENTS_DIR / record["filename"]
        if record["file_format"] == "DOCX":
            build_docx(record, output)
        else:
            build_pdf(record, output, styles)
    serializable = [{k: v for k, v in record.items() if k != "obligation"} for record in records]
    GROUND_TRUTH_PATH.write_text(json.dumps(serializable, ensure_ascii=False, indent=2), encoding="utf-8")
    build_architecture_svg(DATASET_DIR / "architecture.svg")
    print(f"Created {len(records)} documents in {DOCUMENTS_DIR}")


if __name__ == "__main__":
    main()

