from __future__ import annotations

import argparse
import csv
import json
import re
from dataclasses import asdict, dataclass
from datetime import date
from pathlib import Path

from docx import Document
from pypdf import PdfReader


RUSSIAN_MONTHS = {
    "января": 1,
    "февраля": 2,
    "марта": 3,
    "апреля": 4,
    "мая": 5,
    "июня": 6,
    "июля": 7,
    "августа": 8,
    "сентября": 9,
    "октября": 10,
    "ноября": 11,
    "декабря": 12,
}


@dataclass
class ContractFields:
    contract_number: str | None
    contract_date: str | None
    customer: str | None
    contractor: str | None
    amount_rub: int | None


def read_docx(path: Path) -> str:
    """Return text from paragraphs and table cells in their document order groups."""
    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells if cell.text.strip())
    return "\n".join(parts)


def read_pdf(path: Path) -> str:
    """Return text from a PDF that already has a text layer."""
    reader = PdfReader(path)
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def read_document(path: str | Path) -> str:
    path = Path(path)
    if path.suffix.lower() == ".docx":
        return read_docx(path)
    if path.suffix.lower() == ".pdf":
        return read_pdf(path)
    raise ValueError(f"Неподдерживаемый формат: {path.suffix}")


def normalize_text(text: str) -> str:
    lines = [" ".join(line.replace("\xa0", " ").split()) for line in text.splitlines()]
    return "\n".join(line for line in lines if line)


def first_match(patterns: list[str], text: str) -> str | None:
    for pattern in patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE | re.MULTILINE | re.DOTALL)
        if match:
            return " ".join(match.group(1).strip(" \t\n,:;").split())
    return None


def extract_number(text: str) -> str | None:
    return first_match(
        [
            r"(?:договор|контракт)\s*№\s*([A-ZА-ЯЁ0-9][A-ZА-ЯЁ0-9./-]*)",
            r"номер\s+документа\s*\n\s*([A-ZА-ЯЁ0-9][A-ZА-ЯЁ0-9./-]*)",
            r"№\s*([A-ZА-ЯЁ0-9][A-ZА-ЯЁ0-9./-]*)\s+от\b",
        ],
        text,
    )


def extract_date(text: str) -> str | None:
    numeric = first_match(
        [
            r"дата\s+заключения\s*\n\s*(\d{1,2}[./-]\d{1,2}[./-]\d{4})",
            r"(?:\bот\b|новосибирск,)\s*[«\"]?(\d{1,2}[./-]\d{1,2}[./-]\d{4})",
        ],
        text,
    )
    if numeric:
        day, month, year = (int(part) for part in re.split(r"[./-]", numeric))
        return date(year, month, day).isoformat()

    words = re.search(
        r"[«\"]?(\d{1,2})[»\"]?\s+"
        r"(января|февраля|марта|апреля|мая|июня|июля|августа|сентября|октября|ноября|декабря)"
        r"\s+(\d{4})",
        text,
        flags=re.IGNORECASE,
    )
    if words:
        day = int(words.group(1))
        month = RUSSIAN_MONTHS[words.group(2).lower()]
        year = int(words.group(3))
        return date(year, month, day).isoformat()
    return None


def clean_party(value: str | None) -> str | None:
    if value is None:
        return None
    return value.strip().rstrip(". ")


def extract_parties(text: str) -> tuple[str | None, str | None]:
    # Шаблон с отдельными строками и таблицей реквизитов.
    table = re.search(
        r"^Заказчик\s*\n([^\n]+)\s*\nИсполнитель\s*\n([^\n]+)\s*\nСумма договора$",
        text,
        flags=re.IGNORECASE | re.MULTILINE,
    )
    if table:
        return clean_party(table.group(1)), clean_party(table.group(2))

    # Вводный юридический абзац: «..., именуемое ... Заказчик».
    prose = re.search(
        r"^([^\n]+?),\s*именуем\w*\s+в\s+дальнейшем\s+«Заказчик»,\s+и\s+"
        r"(.+?),\s*именуем\w*\s+в\s+дальнейшем\s+«Исполнитель»",
        text,
        flags=re.IGNORECASE | re.MULTILINE | re.DOTALL,
    )
    if prose:
        return clean_party(prose.group(1)), clean_party(prose.group(2).replace("\n", " "))

    # Формат «Сторона 1 / Сторона 2».
    sides = re.search(
        r"Сторона\s*1\s*\(Заказчик\)\s*:\s*(.+?)\.\s*"
        r"Сторона\s*2\s*\(Исполнитель\)\s*:\s*(.+?)\.",
        text,
        flags=re.IGNORECASE | re.DOTALL,
    )
    if sides:
        return clean_party(sides.group(1)), clean_party(sides.group(2))

    # Формат с явными подписями «ИСПОЛНИТЕЛЬ:» и «ЗАКАЗЧИК:».
    customer = first_match([r"^ЗАКАЗЧИК\s*:\s*([^\n]+)$"], text)
    contractor = first_match([r"^ИСПОЛНИТЕЛЬ\s*:\s*([^\n]+)$"], text)
    return clean_party(customer), clean_party(contractor)


def extract_amount(text: str) -> int | None:
    raw = first_match(
        [
            r"сумма\s+договора\s*:\s*([\d\s]+(?:[.,]\d{2})?)",
            r"(?:составляет|в\s+размере)\s+([\d\s]+(?:[.,]\d{2})?)",
        ],
        text,
    )
    if raw is None:
        return None
    number = raw.replace(" ", "").replace(",", ".")
    return round(float(number))


def parse_contract(path: str | Path) -> ContractFields:
    text = normalize_text(read_document(path))
    customer, contractor = extract_parties(text)
    return ContractFields(
        contract_number=extract_number(text),
        contract_date=extract_date(text),
        customer=customer,
        contractor=contractor,
        amount_rub=extract_amount(text),
    )


def evaluate_dataset(dataset_dir: str | Path) -> tuple[list[dict], dict]:
    dataset_dir = Path(dataset_dir)
    labels = json.loads((dataset_dir / "ground_truth.json").read_text(encoding="utf-8"))
    field_names = ["contract_number", "contract_date", "customer", "contractor", "amount_rub"]
    rows = []
    correct = {field: 0 for field in field_names}

    for expected in labels:
        actual = asdict(parse_contract(dataset_dir / "documents" / expected["filename"]))
        row = {"filename": expected["filename"]}
        all_correct = True
        for field in field_names:
            matches = actual[field] == expected[field]
            correct[field] += int(matches)
            all_correct = all_correct and matches
            row[f"expected_{field}"] = expected[field]
            row[f"predicted_{field}"] = actual[field]
            row[f"correct_{field}"] = matches
        row["all_fields_correct"] = all_correct
        rows.append(row)

    total = len(labels)
    metrics = {
        "documents": total,
        "exact_document_accuracy": sum(row["all_fields_correct"] for row in rows) / total,
        "field_accuracy": {field: correct[field] / total for field in field_names},
    }
    return rows, metrics


def save_results(rows: list[dict], path: str | Path) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    parser = argparse.ArgumentParser(description="Baseline-парсер синтетических договоров")
    parser.add_argument(
        "--dataset",
        type=Path,
        default=Path(__file__).parent / "contract_dataset_v1",
        help="Путь к папке датасета",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(__file__).parent / "baseline_results.csv",
        help="Куда сохранить подробные результаты",
    )
    args = parser.parse_args()

    rows, metrics = evaluate_dataset(args.dataset)
    save_results(rows, args.output)
    print(json.dumps(metrics, ensure_ascii=False, indent=2))
    print(f"Подробные результаты: {args.output}")


if __name__ == "__main__":
    main()

